#!/usr/bin/env python3
"""
build_tracker_csv.py — make data/energy_monthly.csv for the payback panel.

Inputs
  --billing  Energy usage workbook with an "Energy" sheet of E.ON billing rows:
             Date Period | Days | E.Usage (kWh) | E.Price | E.Standing Charge (rate) |
             E.Standing Charge (£) | E.Cost ...
  --sigen    Optional Sigenergy hourly export (stationData): time, grid import,
             grid export, load, solar, battery charge, battery discharge.
             The app's header row omits the time column, so values sit one column
             right of their labels; this script reads them by position.
  --ha-start First day Home Assistant recorded (default 2026-04-09). Sigenergy rows
             from that day on are skipped so nothing is counted twice.

Output columns
  label, import_kwh, export_kwh, load_kwh, import_cost, export_credit, standing_charge

Months with Sigenergy data use its import/export/load kWh, priced at that month's
billed unit rate (E.Price + 5% VAT). Other months use the billing rows as they are.
Export credit is left at 0: the dashboard accrues it from export_kwh at the tariff's
export rate.

    python3 tools/build_tracker_csv.py --billing Energy_usage_2024-.xlsx \
        --sigen sigen-hourly.xlsx --out data/energy_monthly.csv
"""
import argparse
import csv
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook

VAT = 1.05


def read_billing(path):
    ws = load_workbook(path, data_only=True)["Energy"]
    months = defaultdict(lambda: {"kwh": 0.0, "cost": 0.0, "standing": 0.0, "price_kwh": 0.0})
    for r in list(ws.iter_rows(values_only=True))[1:]:
        d, _days, kwh, price, _sc_rate, sc_total, cost = r[:7]
        if not isinstance(d, datetime) or not kwh:
            continue
        m = months[d.strftime("%Y-%m")]
        m["kwh"] += kwh
        m["cost"] += cost or 0
        m["standing"] += (sc_total or 0) * VAT
        m["price_kwh"] += (price or 0) * kwh     # kWh-weighted unit price
    for m in months.values():
        m["price"] = m["price_kwh"] / m["kwh"] if m["kwh"] else 0
    return months


def read_sigen(path, ha_start):
    ws = load_workbook(path, data_only=True).active
    months = defaultdict(lambda: {"imp": 0.0, "exp": 0.0, "load": 0.0})
    for r in ws.iter_rows(min_row=2, values_only=True):
        t = r[0]
        if not t or str(t).lower() == "total":
            continue
        ts = t if isinstance(t, datetime) else datetime.fromisoformat(str(t))
        if ts.date() >= ha_start:
            continue
        imp, exp, load = (max(0.0, float(v or 0)) for v in r[1:4])
        m = months[ts.strftime("%Y-%m")]
        m["imp"] += imp
        m["exp"] += exp
        m["load"] += load
    return months


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--billing", required=True)
    ap.add_argument("--sigen")
    ap.add_argument("--ha-start", default="2026-04-09")
    ap.add_argument("--out", default="data/energy_monthly.csv")
    a = ap.parse_args()

    ha_start = datetime.fromisoformat(a.ha_start).date()
    bill = read_billing(a.billing)
    sigen = read_sigen(a.sigen, ha_start) if a.sigen else {}
    last = ha_start.strftime("%Y-%m")

    rows = []
    for key in sorted(set(bill) | set(sigen)):
        if key > last:
            continue
        b = bill.get(key)
        s = sigen.get(key)
        standing = round(b["standing"], 2) if b else 0
        if s:
            price = b["price"] if b else 0
            rows.append({"label": key, "import_kwh": round(s["imp"], 1), "export_kwh": round(s["exp"], 1),
                         "load_kwh": round(s["load"], 1), "import_cost": round(s["imp"] * price * VAT, 2),
                         "export_credit": 0, "standing_charge": standing})
        elif key < last:
            # Before the system: everything the house used came from the grid
            rows.append({"label": key, "import_kwh": round(b["kwh"], 1), "export_kwh": 0,
                         "load_kwh": round(b["kwh"], 1), "import_cost": round(b["cost"] - b["standing"], 2),
                         "export_credit": 0, "standing_charge": standing})

    with open(a.out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["label", "import_kwh", "export_kwh", "load_kwh", "import_cost", "export_credit", "standing_charge"])
        w.writeheader()
        w.writerows(rows)
    print(f"Wrote {len(rows)} months to {a.out}")


if __name__ == "__main__":
    main()
